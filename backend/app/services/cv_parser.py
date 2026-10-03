import io

from docx import Document
from pypdf import PdfReader


class CVError(ValueError):
    pass


def extract_text(file_name: str, data: bytes) -> str:
    name = file_name.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((p.extract_text() or "") for p in reader.pages).strip()
        if len(text) < 200:
            raise CVError("Scanned PDF not supported. Upload a text PDF or DOCX.")
        return text
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts).strip()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore").strip()
    raise CVError("Unsupported file. Use PDF, DOCX, TXT or MD.")
